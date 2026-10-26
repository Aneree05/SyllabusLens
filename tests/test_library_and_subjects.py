"""Library ingestion, validation, extraction, duplicates, and subject management. Synthetic data."""
import io

import pytest

import config
from services.errors import (
    ConfirmationError, DocumentValidationError, DuplicateDocumentError, NotFoundError,
)


def _pdf(pages):
    rl = pytest.importorskip("reportlab.pdfgen.canvas")
    buf = io.BytesIO(); c = rl.Canvas(buf)
    for t in pages:
        c.drawString(72, 720, t); c.showPage()
    c.save(); return buf.getvalue()


def _pptx(slides):
    from pptx import Presentation
    prs = Presentation()
    for title, body in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = title; s.placeholders[1].text = body
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


@pytest.fixture
def subj(svc):
    return svc.subjects.create("Test Subject")


def test_pdf_extraction_keeps_page_references(svc, subj):
    r = svc.library.ingest(subj.id, "n.pdf", _pdf(["Tokenization splits text into tokens.", "Parsing builds trees."]), doc_kind="lecture_notes")
    secs = svc.library.sections_of(r.document.id)
    assert r.document.status == "processed" and r.document.page_count == 2
    assert [s["page_or_slide"] for s in secs] == ["Page 1", "Page 2"]
    assert "Tokenization" in secs[0]["text"]


def test_pptx_extraction_keeps_slide_references(svc, subj):
    r = svc.library.ingest(subj.id, "s.pptx", _pptx([("Embeddings", "word2vec skip gram"), ("Parsing", "cyk algorithm")]))
    assert r.kind_suggestion[0] == "lecture_slides" and r.document.doc_kind == "lecture_slides"
    secs = svc.library.sections_of(r.document.id)
    assert [s["page_or_slide"] for s in secs] == ["Slide 1", "Slide 2"] and "Embeddings" in secs[0]["heading"]


def test_image_only_pdf_is_flagged_not_silently_empty(svc, subj):
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(72, 72); b = io.BytesIO(); w.write(b)
    r = svc.library.ingest(subj.id, "scan.pdf", b.getvalue(), doc_kind="lecture_notes")
    assert r.document.status == "no_text" and "OCR" in r.document.status_detail
    assert svc.library.stats(subj.id)["failures"] == 1
    with pytest.raises(DocumentValidationError):
        svc.retrieval.build_note_corpus(subj.id, [r.document.id])


@pytest.mark.parametrize("name,data,msg", [
    ("a.exe", b"MZ....", "unsupported"), ("empty.txt", b"", "empty"),
    ("fake.pdf", b"this is not a pdf", "not a valid PDF"), ("fake.pptx", b"hello", "not a valid .pptx"),
    ("bin.txt", b"abc\x00\x01\x02def", "binary"), ("noext", b"hello", "unsupported"),
])
def test_malformed_uploads_rejected_with_actionable_message(svc, subj, name, data, msg):
    with pytest.raises(DocumentValidationError, match=msg):
        svc.library.ingest(subj.id, name, data)


def test_size_limit(svc, subj, monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 100)
    with pytest.raises(DocumentValidationError, match="upload limit"):
        svc.library.ingest(subj.id, "big.txt", b"x" * 101)


def test_path_traversal_filename_is_neutralised(svc, subj):
    r = svc.library.ingest(subj.id, "../../etc/evil.txt", b"# H\nhello world text", doc_kind="lecture_notes")
    assert r.document.filename == "evil.txt"
    assert (svc.files.root / subj.id / r.document.stored_name).exists()
    assert not (svc.data_dir.parent / "etc").exists()


def test_duplicate_in_same_subject_blocked_unless_allowed(svc, subj):
    a = svc.library.ingest(subj.id, "a.txt", b"# H\nsame bytes here", doc_kind="lecture_notes").document
    with pytest.raises(DuplicateDocumentError) as e:
        svc.library.ingest(subj.id, "copy.txt", b"# H\nsame bytes here", doc_kind="lecture_notes")
    assert e.value.existing_id == a.id
    assert svc.library.ingest(subj.id, "copy.txt", b"# H\nsame bytes here", doc_kind="lecture_notes", allow_duplicate=True)


def test_same_text_different_bytes_warns(svc, subj):
    svc.library.ingest(subj.id, "a.txt", b"# H\nsome text here", doc_kind="lecture_notes")
    r = svc.library.ingest(subj.id, "b.txt", b"# H\nsome  text   here\n", doc_kind="lecture_notes")
    assert r.same_subject_text_duplicates and r.warnings


def test_search_filter_sort_and_content_search(svc):
    a = svc.subjects.create("A"); b = svc.subjects.create("B")
    svc.library.ingest(a.id, "alpha.txt", b"# T\nquantum entanglement", doc_kind="lecture_notes", tags=["physics"], title="Alpha notes")
    svc.library.ingest(a.id, "beta.txt", b"# T\nbaking bread", doc_kind="assignment")
    svc.library.ingest(b.id, "gamma.txt", b"# T\nquantum tunnelling", doc_kind="lecture_notes")
    L = svc.library
    assert [d.filename for d in L.list_documents(a.id, kind="assignment")] == ["beta.txt"]
    assert {d.filename for d in L.list_documents(a.id, query="quantum", search_content=True)} == {"alpha.txt"}
    assert L.list_documents(a.id, query="quantum") == []                       # title/filename only by default
    assert {d.filename for d in L.list_documents(None, query="quantum", search_content=True)} == {"alpha.txt", "gamma.txt"}
    assert [d.filename for d in L.list_documents(a.id, query="physics")] == ["alpha.txt"]
    assert [d.filename for d in L.list_documents(a.id, sort="filename")] == ["alpha.txt", "beta.txt"]
    assert L.list_documents(a.id, query="100%_") == []                          # LIKE wildcards are escaped
    assert len(L.list_documents(a.id, limit=1)) == 1 and len(L.list_documents(a.id, limit=1, offset=1)) == 1


def test_rename_replace_reclassify_delete_bulk(svc, subj):
    L = svc.library
    d = L.ingest(subj.id, "n.txt", b"# One\nalpha beta gamma\n# Two\ndelta epsilon", doc_kind="lecture_notes").document
    assert L.rename(d.id, "  Week 1 notes ").title == "Week 1 notes"
    with pytest.raises(DocumentValidationError):
        L.rename(d.id, "   ")
    r = L.replace(d.id, "n2.txt", b"# Only\nzeta eta theta")
    assert r.id == d.id and r.file_hash != d.file_hash and len(L.sections_of(d.id)) == 1
    assert L.set_kind(d.id, "question_paper").is_question_paper and L.sections_of(d.id) == []   # papers hold no note sections
    assert len(L.sections_of(L.set_kind(d.id, "lecture_notes").id)) == 1                        # rebuilt from stored file
    assert L.bulk_update([d.id], add_tags=["x", "y"]) == 1 and L.get(d.id).tags == ["x", "y"]
    L.delete(d.id)
    with pytest.raises(NotFoundError):
        L.get(d.id)
    assert not any(svc.files.root.rglob(d.stored_name))


def test_question_corrections_persist_and_feed_analysis(svc, world):
    from nlp.question_segmentation import QuestionItem
    L = svc.library
    p = world["nlp_p1"]
    auto = L.questions_of(p.id)
    L.save_question_corrections(p.id, [QuestionItem("Q1", "x", "Explain word embeddings and word2vec", "5")])
    assert [q.clean_text for q in L.questions_of(p.id)] == ["Explain word embeddings and word2vec"]
    out = svc.runner.run(world["nlp"].id, [p.id], save=False)
    assert "Word embeddings" in out.record["questions"][0]["matches"][0]["heading"] and len(out.record["questions"]) == 1
    with pytest.raises(DocumentValidationError):
        L.save_question_corrections(p.id, [])
    L.clear_question_corrections(p.id)
    assert len(L.questions_of(p.id)) == len(auto)


def test_subject_crud_and_validation(svc):
    s = svc.subjects.create("  Machine   Learning ", description="d", course_code="CS501", semester="7", topics=["SVM", "svm", " ", "Kernels"])
    assert s.name == "Machine Learning" and s.topics == ["SVM", "Kernels"]
    with pytest.raises(DocumentValidationError):
        svc.subjects.create("machine learning")                                   # case-insensitive uniqueness
    with pytest.raises(DocumentValidationError):
        svc.subjects.create("   ")
    u = svc.subjects.update(s.id, name="ML", topics=["Trees"])
    assert (u.name, u.topics, u.course_code) == ("ML", ["Trees"], "CS501")
    with pytest.raises(NotFoundError):
        svc.subjects.get("sub_nope")


def test_subject_delete_requires_exact_name_and_removes_everything(svc, world):
    sid = world["os"].id
    stored = [svc.files.root / sid / d.stored_name for d in world["os_docs"]]
    assert all(p.exists() for p in stored)
    with pytest.raises(ConfirmationError):
        svc.subjects.delete(sid, "wrong name")
    assert svc.subjects.exists(sid)
    svc.runner.run(sid, [world["os_p1"].id])
    removed = svc.subjects.delete(sid, "Operating Systems")
    assert removed["documents"] == 2 and removed["analyses"] == 1
    assert not svc.subjects.exists(sid) and not any(p.exists() for p in stored)
    assert svc.library.list_documents(sid) == []
    assert svc.subjects.exists(world["nlp"].id) and svc.library.list_documents(world["nlp"].id)   # NLP untouched


def test_unknown_subject_upload_rejected(svc):
    with pytest.raises(NotFoundError):
        svc.library.ingest("sub_missing", "a.txt", b"hello world")
