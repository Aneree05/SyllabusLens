"""
Real environment diagnostics (nothing here is hardcoded to "OK").

Windows Application Control / WDAC / AppLocker can block native extension modules
(.pyd/.dll) such as SciPy's. Python then raises ImportError/OSError mentioning the policy
(commonly WinError 4551). This module detects and REPORTS that condition with guidance. It
never tries to bypass or disable the policy.
"""
from __future__ import annotations

import importlib
import io
import platform
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

REQUIRED = [("numpy", "numpy"), ("scipy", "scipy"), ("sklearn", "scikit-learn"), ("pandas", "pandas"),
            ("plotly", "plotly"), ("pypdf", "pypdf"), ("pptx", "python-pptx"), ("streamlit", "streamlit")]

_POLICY_MARKERS = ("application control policy", "blocked this file", "winerror 4551", "4551",
                   "device guard", "applocker", "code integrity")

POLICY_GUIDANCE = (
    "A Windows security policy (Application Control / WDAC / AppLocker) appears to have blocked a native "
    "library. SyllabusLens does not attempt to bypass it. Options: (1) ask your IT administrator to allow the "
    "blocked file shown in the error; (2) use a Python distribution that is already approved on this machine "
    "(e.g. the python.org installer or Microsoft Store build) and reinstall dependencies into a fresh virtual "
    "environment; (3) run SyllabusLens on a machine or container where the policy does not apply. "
    "Reinstalling the same unsigned wheel will not change the policy decision."
)


@dataclass
class Check:
    name: str
    ok: bool
    detail: str = ""
    version: str = ""
    blocked_by_policy: bool = False


@dataclass
class DiagnosticsReport:
    python: str
    platform: str
    checks: List[Check] = field(default_factory=list)
    policy_blocked: bool = False
    guidance: str = ""

    @property
    def ok(self) -> bool:
        return all(c.ok for c in self.checks)

    def failed(self) -> List[Check]:
        return [c for c in self.checks if not c.ok]


def _looks_like_policy_block(exc: BaseException) -> bool:
    text = f"{exc} {getattr(exc, 'winerror', '')}".lower()
    return any(m in text for m in _POLICY_MARKERS)


def check_import(module: str, label: str) -> Check:
    try:
        mod = importlib.import_module(module)
        return Check(label, True, "imported", getattr(mod, "__version__", ""))
    except (ImportError, OSError) as exc:
        blocked = _looks_like_policy_block(exc)
        return Check(label, False, f"{type(exc).__name__}: {exc}", blocked_by_policy=blocked)
    except Exception as exc:  # pragma: no cover - defensive
        return Check(label, False, f"{type(exc).__name__}: {exc}")


def check_functionality() -> List[Check]:
    """Exercise the capabilities the app actually needs, not just imports."""
    out: List[Check] = []
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity
        m = TfidfVectorizer().fit_transform(["alpha beta", "beta gamma"])
        v = float(cosine_similarity(m)[0, 1])
        out.append(Check("TF-IDF + cosine similarity", 0.0 < v < 1.0, f"sample similarity {v:.3f}"))
    except Exception as exc:
        out.append(Check("TF-IDF + cosine similarity", False, f"{type(exc).__name__}: {exc}",
                         blocked_by_policy=_looks_like_policy_block(exc)))
    try:
        from pptx import Presentation
        from services.document_service import DocumentService
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Diagnostic"
        s.placeholders[1].text = "tokenization probe"
        buf = io.BytesIO(); prs.save(buf)
        ex = DocumentService().extract_from_bytes(buf.getvalue(), "probe.pptx")
        out.append(Check("PPTX extraction", "tokenization" in ex.full_text, "round-trip probe"))
    except Exception as exc:
        out.append(Check("PPTX extraction", False, f"{type(exc).__name__}: {exc}"))
    try:
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(); w.add_blank_page(72, 72); buf = io.BytesIO(); w.write(buf); buf.seek(0)
        out.append(Check("PDF parsing", len(PdfReader(buf).pages) == 1, "blank-page round trip (text layer not tested)"))
    except Exception as exc:
        out.append(Check("PDF parsing", False, f"{type(exc).__name__}: {exc}"))
    return out


def check_storage(data_dir: Path) -> Check:
    try:
        data_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=data_dir, delete=True):
            pass
        return Check("Data directory writable", True, str(data_dir))
    except Exception as exc:
        return Check("Data directory writable", False, f"{data_dir}: {exc}")


def run_diagnostics(data_dir: Optional[Path] = None) -> DiagnosticsReport:
    import config
    rep = DiagnosticsReport(python=sys.version.split()[0], platform=platform.platform())
    for module, label in REQUIRED:
        rep.checks.append(check_import(module, label))
    if all(c.ok for c in rep.checks):  # functional probes only make sense if imports worked
        rep.checks.extend(check_functionality())
    rep.checks.append(check_storage(Path(data_dir or config.DATA_DIR)))
    rep.policy_blocked = any(c.blocked_by_policy for c in rep.checks)
    if rep.policy_blocked:
        rep.guidance = POLICY_GUIDANCE
    return rep
