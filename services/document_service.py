"""
Document Extraction and Management Service for SyllabusLens.
Handles PDF (pypdf), PPTX (python-pptx), and TXT text extraction,
with page/slide tracking, scanned PDF detection, and error handling.
"""

import io
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional, Tuple
from pypdf import PdfReader
from pptx import Presentation
from nlp.document_segmentation import DocumentSegmenter, NoteSection


@dataclass
class ExtractedDocument:
    """Represents an uploaded document and its extracted contents."""
    filename: str
    doc_type: str   # "pdf", "pptx", "txt"
    file_size_bytes: int
    page_or_slide_count: int
    full_text: str
    is_likely_scanned: bool = False
    error_message: Optional[str] = None
    sections: List[NoteSection] = field(default_factory=list)
    page_texts: List[str] = field(default_factory=list)
    slides_data: List[Dict[str, Any]] = field(default_factory=list)


class DocumentService:
    """
    Coordinates file parsing, text extraction, validation,
    and segmentation into referenced note sections.
    """

    def __init__(self, segmenter: Optional[DocumentSegmenter] = None):
        self.segmenter = segmenter or DocumentSegmenter()

    def extract_from_bytes(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        """
        Extracts content from in-memory file bytes based on file extension.
        """
        ext = filename.lower().split('.')[-1] if '.' in filename else ""
        file_size = len(file_bytes)

        if ext == "pdf":
            return self._extract_pdf(file_bytes, filename, file_size)
        elif ext == "pptx":
            return self._extract_pptx(file_bytes, filename, file_size)
        elif ext in ("txt", "md"):
            return self._extract_txt(file_bytes, filename, file_size)
        else:
            return ExtractedDocument(
                filename=filename,
                doc_type=ext,
                file_size_bytes=file_size,
                page_or_slide_count=0,
                full_text="",
                error_message=f"Unsupported file format: .{ext}. Supported formats: PDF, PPTX, TXT."
            )

    def _extract_pdf(self, file_bytes: bytes, filename: str, file_size: int) -> ExtractedDocument:
        """Extracts text page-by-page using pypdf."""
        try:
            stream = io.BytesIO(file_bytes)
            reader = PdfReader(stream)
            num_pages = len(reader.pages)
            page_texts: List[str] = []

            for p_idx, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                page_texts.append(extracted)

            full_text = "\n\n".join(page_texts).strip()

            # Detect scanned PDF: if document has pages but total characters are extremely low
            total_chars = len(full_text.strip())
            is_scanned = (num_pages > 0 and (total_chars / max(1, num_pages)) < 30)

            # Segment into sections
            sections = self.segmenter.segment_pdf_pages(filename, page_texts)

            return ExtractedDocument(
                filename=filename,
                doc_type="pdf",
                file_size_bytes=file_size,
                page_or_slide_count=num_pages,
                full_text=full_text,
                is_likely_scanned=is_scanned,
                page_texts=page_texts,
                sections=sections
            )
        except Exception as exc:
            return ExtractedDocument(
                filename=filename,
                doc_type="pdf",
                file_size_bytes=file_size,
                page_or_slide_count=0,
                full_text="",
                error_message=f"PDF extraction error: {str(exc)}"
            )

    def _extract_pptx(self, file_bytes: bytes, filename: str, file_size: int) -> ExtractedDocument:
        """Extracts text slide-by-slide using python-pptx."""
        try:
            stream = io.BytesIO(file_bytes)
            prs = Presentation(stream)
            num_slides = len(prs.slides)
            slides_data: List[Dict[str, Any]] = []
            full_text_parts: List[str] = []

            for s_idx, slide in enumerate(prs.slides, 1):
                slide_title = ""
                slide_texts = []

                if slide.shapes.title and slide.shapes.title.text:
                    slide_title = slide.shapes.title.text.strip()

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            p_text = paragraph.text.strip()
                            if p_text and p_text != slide_title:
                                slide_texts.append(p_text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_texts:
                                slide_texts.append(" | ".join(row_texts))

                body_text = "\n".join(slide_texts).strip()
                slide_full = f"{slide_title}\n{body_text}".strip()
                full_text_parts.append(slide_full)

                slides_data.append({
                    "slide_num": s_idx,
                    "title": slide_title,
                    "text": body_text or slide_title
                })

            full_text = "\n\n".join(full_text_parts).strip()
            sections = self.segmenter.segment_pptx_slides(filename, slides_data)

            return ExtractedDocument(
                filename=filename,
                doc_type="pptx",
                file_size_bytes=file_size,
                page_or_slide_count=num_slides,
                full_text=full_text,
                slides_data=slides_data,
                sections=sections
            )
        except Exception as exc:
            return ExtractedDocument(
                filename=filename,
                doc_type="pptx",
                file_size_bytes=file_size,
                page_or_slide_count=0,
                full_text="",
                error_message=f"PowerPoint extraction error: {str(exc)}"
            )

    def _extract_txt(self, file_bytes: bytes, filename: str, file_size: int) -> ExtractedDocument:
        """Extracts text from TXT or MD files."""
        text = ""
        # Try UTF-8 first, fallback to latin-1
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = file_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        sections = self.segmenter.segment_txt_document(filename, text)

        return ExtractedDocument(
            filename=filename,
            doc_type="txt",
            file_size_bytes=file_size,
            page_or_slide_count=len(sections) or 1,
            full_text=text,
            sections=sections
        )
